import base64
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from pypdf import PdfReader

from job_agent.models.job_record import JobRecordInput
from job_agent.models.profile import ClaimStatus, EvidenceFact, Experience, ExperienceKind, Skill
from job_agent.services.application_pack import (
    build_application_pack,
    resolve_resume_bundle,
)
from job_agent.services.job_repository import JobRepository
from job_agent.services.local_matcher import match_job_locally, structure_job_locally
from job_agent.services.portable_resume import (
    PortableResumeError,
    build_portable_resume_draft,
    find_latest_resume_manifest,
    find_profile_photo,
    read_resume_manifest,
    resume_artifact_from_manifest,
)
from job_agent.services.tailored_resume import approve_resume_visual_review
from tests.helpers import sample_profile


JD = "负责 SQL、Tableau 数据分析、周报与运营协作；要求本科，每周到岗四天。"

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)


class PortableResumeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.profile = sample_profile()
        repository = JobRepository(self.root / "jobs.sqlite3")
        record = repository.upsert_job(
            JobRecordInput(
                company="示例数据",
                title="数据运营实习生",
                jd_text=JD,
                source="test",
                source_url="https://example.com/job/1",
                location="上海",
            )
        )
        self.job = repository.get_job(record.job_id)
        structured = structure_job_locally(
            JD,
            company=self.job.company,
            title=self.job.title,
            location=self.job.location,
            source="test",
            source_url="https://example.com/job/1",
        )
        self.result = match_job_locally(self.profile, structured)
        self.applications_dir = self.root / "applications"
        self.output = self.applications_dir / "job-1-draft"

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_builds_single_page_evidence_locked_docx_and_pdf(self) -> None:
        files = build_portable_resume_draft(
            self.profile,
            self.job,
            self.result,
            self.output,
        )
        content = json.loads(files.content_json.read_text(encoding="utf-8"))
        manifest = json.loads(files.manifest.read_text(encoding="utf-8"))
        fact_ids = content["truthfulness"]["confirmed_fact_ids"]
        source_statements = {
            fact.id: fact.statement
            for experience in self.profile.experiences
            for fact in experience.facts
            if self.profile.is_application_ready(fact.status)
        }
        rendered_bullets = [
            bullet
            for section in content["experience_sections"]
            for entry in section["entries"]
            for bullet in entry["bullets"]
        ]

        self.assertTrue(files.docx.is_file())
        self.assertTrue(files.pdf.is_file())
        self.assertEqual(len(PdfReader(str(files.pdf)).pages), 1)
        self.assertEqual(
            [bullet["text"] for bullet in rendered_bullets],
            [source_statements[fact_id] for fact_id in fact_ids],
        )
        self.assertNotIn("独立搭建 Tableau", files.content_json.read_text(encoding="utf-8"))
        self.assertNotIn("Tableau", content["summary"])
        # 摘要必须是可投递的求职语言，不得包含内部流程提示
        self.assertNotIn("需由本人", content["summary"])
        self.assertNotIn("本地事实", content["summary"])
        self.assertEqual(content["summary"], "")
        # 空话套话不得进入摘要
        self.assertNotIn("期望把上述能力", content["summary"])
        # 匹配分属内部排序数据，不得出现在投递简历正文
        from docx import Document

        document = Document(str(files.docx))
        docx_text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
            p.text for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs
        )
        self.assertNotIn("匹配", docx_text)
        self.assertNotIn("/100", docx_text)
        self.assertNotIn("match_score", json.dumps(content, ensure_ascii=False))
        self.assertNotIn("match_score", json.dumps(manifest, ensure_ascii=False))
        section_titles = [section["title"] for section in content["experience_sections"]]
        self.assertTrue(section_titles)
        self.assertTrue(set(section_titles) <= {"实习经历", "工作经历", "项目经历"})
        labels = {bullet["label"] for bullet in rendered_bullets}
        self.assertTrue(labels <= {"工作内容", "工作业绩", "项目内容", "项目业绩"})
        self.assertEqual(
            manifest["qa"]["star_bullets"],
            sum(1 for bullet in rendered_bullets if bullet["label"].endswith("业绩")),
        )
        # 求职意向行呈现角色与期望城市（目标公司不出现在简历上）
        self.assertIn(f"求职方向：{self.job.title}", docx_text)
        self.assertNotIn("期望城市", docx_text)

    def test_imported_resume_does_not_repeat_facts_or_export_internal_placeholders(self) -> None:
        project = "项目经历：使用 SQL 整理校园报名记录，核对重复报名并形成活动复盘。"
        education = "教育经历：示例大学信息管理本科，2020 年至 2024 年。"
        metadata = "技能：SQL、Excel。目标岗位：数据运营。期望地点：上海。"
        self.profile.education = []
        self.profile.skills = []
        self.profile.experiences = [Experience(
            id="exp-resume-test", kind=ExperienceKind.OTHER, organization="", role="简历原文事实导入",
            facts=[EvidenceFact(id=f"import-{index}", statement=text, status=ClaimStatus.USER_CONFIRMED)
                   for index, text in enumerate([metadata, project, education])]
            + [EvidenceFact(id="pending", statement="独立完成大型商业平台。", status=ClaimStatus.NEEDS_CONFIRMATION)],
        )]
        self.profile.skills = [Skill(id=f"skill-{name}", name=name,
            evidence_fact_ids=["import-0"], status=ClaimStatus.USER_CONFIRMED) for name in ("SQL", "Excel")]
        files = build_portable_resume_draft(self.profile, self.job, self.result, self.output)
        content = json.loads(files.content_json.read_text(encoding="utf-8"))
        text = "".join(page.extract_text() for page in PdfReader(files.pdf).pages)
        from job_agent.services.resume_editor import validate_resume_content
        validate_resume_content(content)
        from job_agent.services.resume_polish import polish_resume_content_locally
        polished = polish_resume_content_locally(content, JD).content
        self.assertEqual(polished["experience_sections"], content["experience_sections"])
        normalized = "".join(text.split())
        self.assertEqual(content["summary"], "")
        self.assertEqual(content["highlights"], [])
        self.assertEqual(content["experience_sections"][0]["title"], "相关经历")
        self.assertEqual(content["experience_sections"][0]["entries"][0]["organization"], "")
        self.assertEqual(content["experience_sections"][0]["entries"][0]["role"], "")
        for source in (project, education):
            self.assertEqual(normalized.count("".join(source.split())), 1)
        self.assertNotIn("机构/项目名称待完善", normalized)
        self.assertNotIn("简历原文事实导入", normalized)
        self.assertNotIn("目标岗位", normalized)
        self.assertNotIn("独立完成大型商业平台", normalized)
        self.assertEqual(set(content["truthfulness"]["confirmed_fact_ids"]), {"import-1", "import-2"})
        self.assertEqual(content["selection"]["metadata_fact_ids"], ["import-0"])
        self.assertIn("import-0", content["selection"]["omitted_fact_ids"])
        self.assertEqual(self.profile.experiences[0].facts[0].statement, metadata)

    def test_completed_education_never_invents_current_student_status(self) -> None:
        self.profile.education[0].end = "2020-06"
        files = build_portable_resume_draft(self.profile, self.job, self.result, self.output)
        content = json.loads(files.content_json.read_text(encoding="utf-8"))
        text = "".join(page.extract_text() for page in PdfReader(files.pdf).pages)
        self.assertNotIn("在读", text)
        self.assertEqual(content["education"][0]["end"], "2020-06")
        self.assertIn("2020.06", text)

    def test_imported_skills_heading_keeps_qualification_and_experience_verbatim(self) -> None:
        from job_agent.services.profile_onboarding import ProfileOnboardingInput, onboard_profile
        from job_agent.services.profile_store import load_profile
        source = "合成人物\n专业技能：持有注册会计师证书，曾负责年度财务审计。\n项目经历：使用 Excel 整理财务记录并核对报表。"
        onboard_profile(ProfileOnboardingInput(display_name="合成人物", resume_filename="resume.txt",
            resume_bytes=source.encode("utf-8"), confirm_truth=True, target_roles=["财务审计"]),
            profile_path=self.root / "private/profile.json", private_dir=self.root / "private")
        profile = load_profile(self.root / "private/profile.json")
        jd = structure_job_locally("财务审计岗位，要求持有注册会计师证书并具有年度财务审计经历。",
            company="合成公司", title="财务审计", location="上海", source="test")
        result = match_job_locally(profile, jd)
        files = build_portable_resume_draft(profile, self.job, result, self.output)
        content = json.loads(files.content_json.read_text(encoding="utf-8"))
        fact = next(fact for exp in profile.experiences for fact in exp.facts if "注册会计师" in fact.statement)
        text = "".join(page.extract_text() for page in PdfReader(files.pdf).pages)
        self.assertIn("".join(fact.statement.split()), "".join(text.split()))
        self.assertIn(fact.id, content["truthfulness"]["confirmed_fact_ids"])
        self.assertNotIn(fact.id, content["selection"]["metadata_fact_ids"])

    def test_other_fact_is_not_duplicated_in_highlights_and_keeps_provenance(self) -> None:
        statement = "获得校级数据分析竞赛优秀奖。"
        self.profile.skills = []
        self.profile.experiences = [Experience(
            id="award", kind=ExperienceKind.OTHER, organization="", role="",
            facts=[EvidenceFact(id="award-1", statement=statement, status=ClaimStatus.USER_CONFIRMED)],
        )]
        files = build_portable_resume_draft(self.profile, self.job, self.result, self.output)
        content = json.loads(files.content_json.read_text(encoding="utf-8"))
        text = "".join(page.extract_text() for page in PdfReader(files.pdf).pages)
        self.assertEqual(content["highlights"], [])
        self.assertEqual("".join(text.split()).count(statement), 1)
        self.assertEqual(content["truthfulness"]["confirmed_fact_ids"], ["award-1"])

    def test_skill_beyond_export_limit_is_not_silently_removed_as_metadata(self) -> None:
        from job_agent.services.portable_resume import build_portable_resume_content
        self.profile.skills = []
        self.profile.experiences = [Experience(id="exp-resume-limit", kind=ExperienceKind.OTHER,
            organization="", role="简历原文事实导入", facts=[
                EvidenceFact(id="skill-source", statement="技能：SQL。", status=ClaimStatus.USER_CONFIRMED),
                EvidenceFact(id="project-source", statement="整理活动报名资料并核对记录。", status=ClaimStatus.USER_CONFIRMED),
            ])]
        self.profile.skills = [Skill(id=f"limit-{index}", name=name,
            evidence_fact_ids=["skill-source"], status=ClaimStatus.USER_CONFIRMED)
            for index, name in enumerate([f"已确认技能{n}" for n in range(12)] + ["SQL"])]
        content = build_portable_resume_content(self.profile, self.job, self.result)
        self.assertNotIn("SQL", content["skills"])
        self.assertNotIn("skill-source", content["selection"]["metadata_fact_ids"])
        self.assertIn("skill-source", content["truthfulness"]["confirmed_fact_ids"])
        self.assertTrue(any(bullet["text"] == "技能：SQL。" for section in content["experience_sections"]
            for entry in section["entries"] for bullet in entry["bullets"]))

    def test_finds_photo_in_assets_fallback_dir(self) -> None:
        # 历史/手动落点：photo/ 不存在时应能找到 assets/ 里的照片
        assets_dir = self.root / "assets"
        assets_dir.mkdir()
        photo = assets_dir / "profile-photo.png"
        photo.write_bytes(_PNG_1PX)
        self.assertEqual(find_profile_photo(self.root), photo)
        # photo/ 存在时优先于 assets/
        photo_dir = self.root / "photo"
        photo_dir.mkdir()
        preferred = photo_dir / "profile-photo.jpg"
        preferred.write_bytes(_PNG_1PX)
        self.assertEqual(find_profile_photo(self.root), preferred)

    def test_embeds_profile_photo_and_reports_in_manifest(self) -> None:
        photo_dir = self.root / "photo"
        photo_dir.mkdir()
        photo = photo_dir / "profile-photo.png"
        photo.write_bytes(_PNG_1PX)

        self.assertEqual(find_profile_photo(self.root), photo)
        files = build_portable_resume_draft(
            self.profile,
            self.job,
            self.result,
            self.output,
            photo_path=photo,
        )
        content = json.loads(files.content_json.read_text(encoding="utf-8"))
        manifest = json.loads(files.manifest.read_text(encoding="utf-8"))

        self.assertEqual(content["person"]["photo_path"], str(photo.resolve()))
        self.assertEqual(manifest["qa"]["embedded_photos"], 1)
        self.assertEqual(manifest["qa"]["photo_requirement"], "embedded")
        self.assertEqual(len(PdfReader(str(files.pdf)).pages), 1)
        with zipfile.ZipFile(files.docx) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
        self.assertEqual(len(media), 1)

    def test_rejects_photo_that_is_missing_or_wrong_type(self) -> None:
        with self.assertRaises(PortableResumeError):
            build_portable_resume_draft(
                self.profile,
                self.job,
                self.result,
                self.output,
                photo_path=self.root / "missing.png",
            )
        bad = self.root / "photo.txt"
        bad.write_text("not a photo", encoding="utf-8")
        with self.assertRaises(PortableResumeError):
            build_portable_resume_draft(
                self.profile,
                self.job,
                self.result,
                self.output,
                photo_path=bad,
            )

    def test_requires_visual_approval_before_application_pack_is_ready(self) -> None:
        files = build_portable_resume_draft(
            self.profile,
            self.job,
            self.result,
            self.output,
        )
        draft = resolve_resume_bundle(
            self.profile,
            self.job,
            self.applications_dir,
        )
        self.assertEqual(draft.status, "needs_review")

        approved = approve_resume_visual_review(files.manifest)
        ready = resolve_resume_bundle(
            self.profile,
            self.job,
            self.applications_dir,
        )
        pack = build_application_pack(self.profile, self.job, self.result, ready)

        self.assertTrue(approved.is_file())
        self.assertEqual(ready.status, "ready")
        self.assertTrue(ready.qa_verified)
        self.assertEqual(pack.resume.status, "ready")
        self.assertEqual(
            find_latest_resume_manifest(self.applications_dir, self.job.job_id),
            approved,
        )
        self.assertEqual(
            resume_artifact_from_manifest(approved, "pdf"),
            files.pdf,
        )

    def test_malformed_sibling_manifest_does_not_block_healthy_job(self) -> None:
        files = build_portable_resume_draft(
            self.profile, self.job, self.result, self.output,
        )
        broken = self.applications_dir / "other-job" / "resume-version-broken.json"
        broken.parent.mkdir(parents=True)
        for payload in (
            {"target": {"job_id": "damaged"}, "qa": {}, "artifacts": {}},
            ["not a manifest"],
            {"target": {"job_id": True}, "qa": {}, "artifacts": {}},
            {"target": {"job_id": self.job.job_id}, "qa": [], "artifacts": {}},
            {"target": {"job_id": self.job.job_id}, "qa": {"pdf_visual_review": []}, "artifacts": {}},
        ):
            broken.write_text(json.dumps(payload), encoding="utf-8")
            self.assertIsNone(read_resume_manifest(broken))
            self.assertEqual(
                find_latest_resume_manifest(self.applications_dir, self.job.job_id),
                files.manifest.resolve(),
            )
            self.assertEqual(resume_artifact_from_manifest(files.manifest, "pdf"), files.pdf)


if __name__ == "__main__":
    unittest.main()
