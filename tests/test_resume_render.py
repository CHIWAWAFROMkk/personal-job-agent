"""声明式简历模板渲染与模板 ID 接入测试。"""
from __future__ import annotations

import pytest

from job_agent.services.resume_editor import (
    ResumeEditorError,
    supported_template_ids,
    validate_resume_content,
)
from job_agent.services.resume_render import (
    TemplateRenderError,
    list_templates,
    load_template,
    render_docx,
)


def _content(template_id: str = "minimal-mono") -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "template_id": template_id,
        "resume_version_id": "resume-test-1",
        "person": {"name": "张三", "phone": "13800000000", "email": "resume@example.com", "photo_path": ""},
        "target": {"job_id": 7, "company": "某公司", "role": "运营", "location": "上海", "match_score": 90},
        "summary": "面向运营岗位的数据驱动型选手。",
        "skills": ["Excel", "SQL", "Python"],
        "experience_sections": [
            {
                "title": "实习与工作经历",
                "entries": [
                    {
                        "organization": "某公司",
                        "role": "运营实习生",
                        "dates": "2025.06-2025.09",
                        "bullets": [
                            {"label": "拉新", "text": "搭建社群裂变链路，30 天拉新 1200 人。", "fact_ids": ["f1"]},
                            {"label": "转化", "text": "优化活动落地页，转化率提升 18%。", "fact_ids": ["f2"]},
                        ],
                    }
                ],
            }
        ],
        "education": [
            {
                "institution": "某大学",
                "major": "市场营销",
                "degree": "本科",
                "start": "2022-09",
                "end": "2026.06",
                "coursework": ["市场调研", "消费者行为"],
            }
        ],
        "truthfulness": {"confirmed_fact_ids": ["f1", "f2"], "rule": "test"},
    }


def test_list_templates_contains_three_specs() -> None:
    templates = list_templates()
    for tid in ("minimal-mono", "classic-centered", "modern-photo"):
        assert tid in templates


def test_load_template_unknown_id_rejected() -> None:
    with pytest.raises(TemplateRenderError):
        load_template("no-such-template")


def test_target_label_excludes_application_metadata():
    from job_agent.services.resume_render import format_resume_target
    assert format_resume_target({"company": "目标公司", "role": "产品运营（校园招聘） (上海)【急招】", "location": "上海"}) == "求职方向：产品运营"
    assert format_resume_target({"role": "数据运营（实习（可转正））"}) == "求职方向：数据运营"


def test_internship_heading_does_not_relabel_employment():
    from job_agent.services.resume_render import resume_section_title
    assert resume_section_title({"title": "工作经历", "entries": [{"role": "运营实习生"}]}) == "实习经历"
    assert resume_section_title({"title": "工作经历", "entries": [{"role": "运营经理"}]}) == "工作经历"


def test_self_evaluation_renders_and_can_be_hidden(tmp_path):
    from docx import Document
    content = _content("mono-photo")
    content["self_evaluation"] = "善于结合数据复盘运营工作。"
    path = tmp_path / "evaluation.docx"
    render_docx(content, load_template("mono-photo"), path)
    paragraphs = [p.text for p in Document(path).paragraphs]
    assert paragraphs[-2:] == ["自我评价", content["self_evaluation"]]
    content["self_evaluation"] = ""
    render_docx(content, load_template("mono-photo"), tmp_path / "empty.docx")
    assert "自我评价" not in [p.text for p in Document(tmp_path / "empty.docx").paragraphs]


def test_supported_template_ids_includes_legacy_and_specs() -> None:
    ids = supported_template_ids()
    assert "portable-evidence-resume-v1" in ids
    assert "minimal-mono" in ids


def test_validate_accepts_declarative_template_id() -> None:
    assert validate_resume_content(_content("modern-photo")) is not None


def test_validate_rejects_unknown_template_id() -> None:
    with pytest.raises(ResumeEditorError):
        validate_resume_content(_content("not-a-template"))


@pytest.mark.parametrize("template_id", ["minimal-mono", "classic-centered", "modern-photo", "mono-photo"])
def test_render_docx_produces_file(tmp_path, template_id: str) -> None:
    spec = load_template(template_id)
    out = tmp_path / f"{template_id}.docx"
    render_docx(_content(template_id), spec, out)
    assert out.exists() and out.stat().st_size > 10_000


def test_render_docx_omits_internal_match_score(tmp_path) -> None:
    """match_score 是内部排序数据，任何模板都不得写进简历正文。"""
    from docx import Document

    spec = load_template("minimal-mono")
    out = tmp_path / "no-score.docx"
    render_docx(_content("minimal-mono"), spec, out)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "匹配" not in text
    assert "/100" not in text
    assert "某公司" in text and "运营" in text


def test_mono_photo_layout_features(tmp_path) -> None:
    """mono-photo（直聘风）：求职意向行、圆点要点、右对齐日期、主修课程。"""
    from docx import Document

    spec = load_template("mono-photo")
    out = tmp_path / "mono-photo.docx"
    render_docx(_content("mono-photo"), spec, out)
    from docx.oxml.ns import qn
    borders = Document(str(out)).element.xpath("//w:pBdr/w:bottom")
    assert len(borders) == 3
    assert all(item.get(qn("w:color")) == "000000" and item.get(qn("w:sz")) == "6" for item in borders)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    # 求职方向只标职位，不含公司、城市或匹配分。
    assert "求职方向：运营" in text
    assert "期望城市" not in text
    assert "匹配" not in text and "某公司 ·" not in text
    # 个人优势与经历要点都使用圆点，不再让每个分组重新显示“1.”。
    assert "个人优势" in text
    assert "• 面向运营岗位的数据驱动型选手。" in text
    assert "• 搭建社群裂变链路" in text
    assert "• 优化活动落地页" in text
    assert not any(paragraph.startswith(("1. ", "2. ")) for paragraph in text.splitlines())
    # 技能区被该模板关闭（能力并入个人优势）
    assert "核心能力" not in text
    # 日期右对齐条目行：机构+职位加粗、日期保留在同一行
    assert "某公司  运营实习生" in text and "2025.06-2025.09" in text
    # 教育行带日期区间与主修课程
    assert "某大学  市场营销  本科" in text
    assert "2022-09 - 2026.06" in text
    assert "主修课程：市场调研、消费者行为。" in text


def test_mono_photo_uses_reference_document_page_dimensions(tmp_path) -> None:
    from docx import Document

    spec = load_template("mono-photo")
    out = tmp_path / "mono-photo-page-size.docx"
    render_docx(_content("mono-photo"), spec, out)
    section = Document(str(out)).sections[0]

    assert section.page_width.mm == pytest.approx(140, abs=0.2)
    assert section.page_height.mm == pytest.approx(204, abs=0.2)


def test_polish_prompt_enforces_star_writing() -> None:
    from job_agent.services.resume_polish import _SYSTEM_PROMPT as polish_prompt

    assert "STAR" in polish_prompt
    assert "强动词" in polish_prompt
    assert "量化" in polish_prompt


def test_import_prompt_enforces_star_writing() -> None:
    from job_agent.services.resume_import import _SYSTEM_PROMPT as import_prompt

    assert "STAR" in import_prompt
