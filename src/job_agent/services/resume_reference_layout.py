"""Classic A4 layout with shared content order for DOCX and PDF."""
from __future__ import annotations

import html
from pathlib import Path

from job_agent.services.resume_render import format_resume_target, resume_section_title


def availability_line(content: dict) -> str:
    data = content.get("availability") or {}
    values = []
    if data.get("earliest_start"):
        start = str(data["earliest_start"])
        values.append("立即到岗" if start in {"立即", "立刻"} else f"可到岗：{start}")
    if data.get("days_per_week"):
        values.append(f"每周{data['days_per_week']}天")
    if data.get("duration_months"):
        values.append(f"可实习{data['duration_months']}个月以上")
    return "  |  ".join(values)


def body_blocks(content: dict):
    if content.get("education"):
        yield "section", "教育背景", ""
        for item in content["education"]:
            yield "entry", "  |  ".join(str(item.get(k) or "") for k in ("institution", "major", "degree") if item.get(k)), "—".join(str(item.get(k) or "").replace("-", ".") for k in ("start", "end") if item.get(k))
            if item.get("coursework"):
                yield "text", "主修课程：" + "、".join(item["coursework"]), ""
    if content.get("summary"):
        yield "section", "个人优势", ""
        yield "text", content["summary"], ""
    for section in content.get("experience_sections", []):
        yield "section", resume_section_title(section), ""
        for entry in section.get("entries", []):
            yield "entry", "  |  ".join(str(entry.get(k) or "") for k in ("organization", "role") if entry.get(k)), str(entry.get("dates") or "")
            if entry.get("context"):
                yield "text", entry["context"], ""
            for bullet in entry.get("bullets", []):
                yield "bullet", str(bullet.get("text") or ""), str(bullet.get("label") or "")
    if content.get("skills") or content.get("highlights"):
        yield "section", "专业技能与其他", ""
        if content.get("skills"):
            yield "text", " · ".join(content["skills"]), ""
        for item in content.get("highlights", []):
            yield "text", item["text"], ""
    if content.get("self_evaluation"):
        yield "section", "自我评价", ""
        yield "text", content["self_evaluation"], ""


def render_reference_docx(content: dict, spec: dict, path: Path, photo: Path | None):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Mm(210), Mm(297)
    margins = spec["page"]["margins_mm"]
    for side, value in margins.items():
        setattr(page, side + "_margin", Mm(value))
    page.header_distance = page.footer_distance = Mm(4)
    body_size = spec["parts"]["bullet"]["size_pt"]
    leading = spec["parts"]["bullet"]["leading_pt"]
    compact = content.get("layout_density") == "compact"
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = spec["fonts"]["latin"]
        style.font.size = Pt(body_size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)

    def run(p, text, size=body_size, bold=False):
        r = p.add_run(text)
        r.font.name = spec["fonts"]["latin"]
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), spec["fonts"]["eastAsia"])
        r.font.size, r.bold = Pt(size), bold
        r.font.color.rgb = RGBColor(0, 0, 0)

    def spacing(p, before=0, after=2):
        f = p.paragraph_format
        f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), Pt(leading)
        f.widow_control = True

    person = content["person"]
    contacts = "  |  ".join(str(person[k]) for k in ("phone", "email", "city") if person.get(k))
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    for column, width in zip(table.columns, (156, 24)):
        column.width = Mm(width)
    for cell, width in zip(table.rows[0].cells, (156, 24)):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        mar = OxmlElement("w:tcMar")
        for side in ("top", "bottom", "left", "right"):
            node = OxmlElement("w:" + side)
            node.set(qn("w:w"), "0")
            node.set(qn("w:type"), "dxa")
            mar.append(node)
        cell._tc.get_or_add_tcPr().append(mar)
    left, right = table.rows[0].cells
    p = left.paragraphs[0]
    p.style = doc.styles["Title"]
    spacing(p, after=5)
    p.paragraph_format.line_spacing = Pt(25)
    run(p, str(person["name"]), 22, True)
    for value, size, bold in ((format_resume_target(content["target"]), 11, True), (contacts, 10, False), (availability_line(content), 10, False)):
        if value:
            p = left.add_paragraph()
            spacing(p, after=3)
            run(p, value, size, bold)
    if photo:
        p = right.paragraphs[0]
        spacing(p, after=0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(str(photo), width=Mm(spec["photo_width_mm"]))
    for kind, text, extra in body_blocks(content):
        p = doc.add_paragraph(style="Heading 1" if kind == "section" else "Normal")
        spacing(p, before=(4 if compact else 6) if kind == "section" else (0 if compact else 1),
                after=(2 if compact else 4) if kind == "section" else (1 if compact else 2))
        if kind == "section":
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = Pt(17)
            run(p, text, 12, True)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in {"val": "single", "sz": "6", "space": "1", "color": "000000"}.items():
                bottom.set(qn("w:" + k), v)
            border.append(bottom)
            p._p.get_or_add_pPr().append(border)
        elif kind == "entry":
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
            run(p, text, body_size, True)
            if extra:
                run(p, "\t" + extra, 9.5)
        elif kind == "bullet":
            p.paragraph_format.left_indent = Mm(3.3)
            p.paragraph_format.first_line_indent = Mm(-3.3)
            p.paragraph_format.keep_together = True
            run(p, "•  ", 9)
            if extra:
                run(p, extra + "：", body_size, True)
            run(p, text)
        else:
            run(p, text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def render_reference_pdf(content: dict, spec: dict, path: Path, photo: Path | None) -> int:
    from pypdf import PdfReader
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, HRFlowable, Image
    from reportlab.lib.utils import ImageReader

    font = "STSong-Light"
    compact = content.get("layout_density") == "compact"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    base = ParagraphStyle("reference-body", fontName=font, fontSize=spec["parts"]["bullet"]["size_pt"], leading=spec["parts"]["bullet"]["leading_pt"], spaceAfter=1 if compact else 2, wordWrap="CJK")
    styles = {"text": base,
              "section": ParagraphStyle("reference-section", parent=base, fontSize=12, leading=17, spaceBefore=4 if compact else 6, spaceAfter=2 if compact else 3, keepWithNext=True),
              "entry": ParagraphStyle("reference-entry", parent=base, keepWithNext=True, spaceBefore=1),
              "bullet": ParagraphStyle("reference-bullet", parent=base, leftIndent=3.3*mm, firstLineIndent=-3.3*mm)}
    person = content["person"]
    contact = "  |  ".join(str(person[k]) for k in ("phone", "email", "city") if person.get(k))
    left = [Paragraph(html.escape(str(person["name"])), ParagraphStyle("reference-name", parent=base, fontSize=22, leading=25, spaceAfter=5))]
    for text in (format_resume_target(content["target"]), contact, availability_line(content)):
        if text:
            left.append(Paragraph(html.escape(text), base))
    right = ""
    if photo:
        w, h = ImageReader(str(photo)).getSize()
        width = spec["photo_width_mm"] * mm
        right = Image(str(photo), width=width, height=width*h/w)
    table = Table([[left, right]], colWidths=[156*mm, 24*mm])
    table.setStyle(TableStyle([("VALIGN", (0,0),(-1,-1),"MIDDLE"), ("ALIGN",(1,0),(1,0),"RIGHT"),
                               ("LEFTPADDING",(0,0),(-1,-1),0), ("RIGHTPADDING",(0,0),(-1,-1),0),
                               ("TOPPADDING",(0,0),(-1,-1),0), ("BOTTOMPADDING",(0,0),(-1,-1),0)]))
    story = [table]
    for kind, text, extra in body_blocks(content):
        rendered = html.escape(text)
        if kind == "bullet":
            rendered = "•  " + (f"<b>{html.escape(extra)}：</b>" if extra else "") + rendered
        elif kind == "entry":
            # A two-column heading can wrap without placing dates over long company names.
            heading = Table([[Paragraph(f"<b>{rendered}</b>", base), Paragraph(html.escape(extra), base)]], colWidths=[143*mm,37*mm])
            heading.setStyle(TableStyle([("LEFTPADDING",(0,0),(-1,-1),0), ("RIGHTPADDING",(0,0),(-1,-1),0), ("TOPPADDING",(0,0),(-1,-1),1), ("BOTTOMPADDING",(0,0),(-1,-1),2)]))
            heading.keepWithNext = True
            story.append(heading)
            continue
        story.append(Paragraph(rendered, styles[kind]))
        if kind == "section":
            line = HRFlowable(width="100%", thickness=.75, color=colors.black, spaceAfter=3)
            line.keepWithNext = True
            story.append(line)
    margins = spec["page"]["margins_mm"]
    SimpleDocTemplate(str(path), pagesize=(210*mm,297*mm), **{side+"Margin": value*mm for side,value in margins.items()}, title="岗位专属简历", author="个人求职 Agent").build(story)
    return len(PdfReader(path).pages)
