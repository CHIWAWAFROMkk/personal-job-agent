from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path


class ResumeReadError(RuntimeError):
    pass


@dataclass(frozen=True)
class ResumeDocument:
    path: Path
    text: str
    sha256: str
    warnings: list[str] = field(default_factory=list)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ResumeReadError("读取 DOCX 需要安装 python-docx。") from exc
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def _read_pdf(path: Path) -> tuple[str, list[str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ResumeReadError("读取 PDF 需要安装 pypdf。") from exc
    reader = PdfReader(path)
    pages: list[str] = []
    empty_pages: list[int] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if not text.strip():
            empty_pages.append(index)
        pages.append(text)
    warnings: list[str] = []
    if empty_pages:
        warnings.append(
            "以下 PDF 页未提取到文字，可能是扫描件: "
            + ", ".join(map(str, empty_pages))
        )
    return "\n\n".join(pages), warnings


def read_resume(path: Path) -> ResumeDocument:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise ResumeReadError(f"简历文件不存在: {path}")
    suffix = path.suffix.casefold()
    warnings: list[str] = []
    try:
        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8-sig")
        elif suffix == ".docx":
            text = _read_docx(path)
        elif suffix == ".pdf":
            text, warnings = _read_pdf(path)
        else:
            raise ResumeReadError(
                f"暂不支持 {suffix or '无扩展名'}；请使用 TXT、MD、DOCX 或 PDF。"
            )
    except OSError as exc:
        raise ResumeReadError(f"无法读取简历: {exc}") from exc
    if not text.strip():
        raise ResumeReadError("简历中没有提取到可读文字。")
    return ResumeDocument(
        path=path,
        text=text.strip(),
        sha256=_sha256(path),
        warnings=warnings,
    )

