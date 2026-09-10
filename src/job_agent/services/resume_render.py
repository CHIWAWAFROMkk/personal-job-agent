"""声明式简历模板渲染（DOCX 优先）。

模板 spec 见 src/job_agent/templates/*.json：
- 版面/页边距/字体/调色板/语义部件样式全部声明式，渲染器唯一入口。
- 现有硬编码模板（portable-evidence-resume-v1 等）不受影响，本模块并行上线。

设计来源：docs/resume-template-system-plan.md（Reactive Resume 设计提炼，MIT）。
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any


class TemplateRenderError(RuntimeError):
    """模板加载或渲染失败。"""


TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "templates"


def format_resume_target(target: dict[str, Any]) -> str:
    """Label the desired role explicitly; application metadata stays off the resume."""
    role = str(target.get("role") or "").strip()
    while True:
        cleaned = re.sub(r"\([^()]*\)|（[^（）]*）|\[[^\[\]]*\]|【[^【】]*】", "", role)
        if cleaned == role:
            break
        role = cleaned
    role = re.sub(r"\s+", " ", role).strip(" ·|/-—")
    return f"求职方向：{role or '待确认'}"


def resume_section_title(section: dict[str, Any]) -> str:
    title = str(section.get("title") or "经历")
    entries = list(section.get("entries") or [])
    if title in {"工作经历", "实习与工作经历"} and entries and all("实习" in str(entry.get("role") or "") for entry in entries):
        return "实习经历"
    return title


def suggest_self_evaluation(content: dict[str, Any]) -> str:
    skills = [str(item) for item in list(content.get("skills") or [])[:4]]
    roles = list(dict.fromkeys(str(entry.get("role") or "") for section in content.get("experience_sections", []) for entry in section.get("entries", []) if entry.get("role")))[:2]
    sentences = []
    if skills:
        sentences.append(f"技能基础包括{'、'.join(skills)}。")
    if roles:
        sentences.append(f"通过{'、'.join(roles)}的实践积累相关经验。")
    return "".join(sentences)[:300]


def load_template(template_id: str) -> dict[str, Any]:
    """按 ID 加载模板 spec 并做最小校验。"""
    path = TEMPLATES_DIR / f"{template_id}.json"
    if not path.exists():
        raise TemplateRenderError(f"模板不存在: {template_id}（{path}）")
    spec = json.loads(path.read_text(encoding="utf-8"))
    for key in ("id", "layout", "page", "fonts", "palette", "parts"):
        if key not in spec:
            raise TemplateRenderError(f"模板 {template_id} 缺少字段: {key}")
    if spec["id"] != template_id:
        raise TemplateRenderError(f"模板文件与 ID 不一致: {template_id}")
    return spec


def list_templates() -> list[str]:
    """列出可用模板 ID。"""
    if not TEMPLATES_DIR.exists():
        return []
    return sorted(p.stem for p in TEMPLATES_DIR.glob("*.json"))


def _hex_to_rgbcolor(value: str):
    from docx.shared import RGBColor

    return RGBColor.from_string(value.lstrip("#").upper())


def _palette_color(spec: dict[str, Any], name: str):
    palette = spec["palette"]
    raw = palette.get(name, name)
    return _hex_to_rgbcolor(str(raw))


def _apply_run_style(run: Any, spec: dict[str, Any], part: dict[str, Any]) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    fonts = spec["fonts"]
    run.font.name = fonts["latin"]
    run._element.rPr.rFonts.set(qn("w:eastAsia"), fonts["eastAsia"])
    run.font.size = Pt(float(part.get("size_pt", 9)))
    run.bold = bool(part.get("bold", False))
    run.font.color.rgb = _palette_color(spec, str(part.get("color", "body")))


def _align(name: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(name, WD_ALIGN_PARAGRAPH.LEFT)


def _page_dimensions_mm(page: dict[str, Any]) -> tuple[float, float]:
    """Resolve the physical page size declared by a template."""
    size = str(page.get("size", "A4")).upper()
    if size == "A4":
        return 210.0, 297.0
    if size == "CUSTOM":
        try:
            width = float(page["width_mm"])
            height = float(page["height_mm"])
        except (KeyError, TypeError, ValueError) as exc:
            raise TemplateRenderError("自定义页面必须提供有效的 width_mm 与 height_mm。") from exc
        if width <= 0 or height <= 0:
            raise TemplateRenderError("自定义页面尺寸必须大于 0。")
        return width, height
    raise TemplateRenderError(f"不支持的页面尺寸: {page.get('size')}")


def _format_age(value: Any) -> str:
    """年龄字段：数字补"岁"，字符串原样，空值丢弃。"""
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)):
        return f"{int(value)}岁"
    return str(value)


def _add_numbered_items(document: Any, spec: dict[str, Any], part: dict[str, Any], items: list[str]) -> None:
    """手工编号列表（1. 2. 3.），每条独立段落、悬挂缩进，编号随条目组自然重新开始。"""
    from docx.shared import Mm, Pt

    indent = float(part.get("indent_mm", 5))
    for index, item in enumerate(items, 1):
        text = item if item.endswith(("。", "；", ";", ".", "！", "？", "!", "?")) else f"{item}。"
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_after = Pt(float(part.get("space_after_pt", 2)))
        if part.get("leading"):
            fmt.line_spacing = float(part["leading"])
        fmt.left_indent = Mm(indent)
        fmt.first_line_indent = Mm(-indent)
        run = paragraph.add_run(f"{index}. {text}")
        _apply_run_style(run, spec, part)


def _add_bulleted_items(document: Any, spec: dict[str, Any], part: dict[str, Any], items: list[str]) -> None:
    """Render compact bullets without Word's restart-prone numbering metadata."""
    from docx.shared import Mm, Pt

    indent = float(part.get("indent_mm", 5))
    for item in items:
        text = item if item.endswith(("。", "；", ";", ".", "！", "？", "!", "?")) else f"{item}。"
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_after = Pt(float(part.get("space_after_pt", 2)))
        if part.get("leading"):
            fmt.line_spacing = float(part["leading"])
        fmt.left_indent = Mm(indent)
        fmt.first_line_indent = Mm(-indent)
        run = paragraph.add_run(f"• {text}")
        _apply_run_style(run, spec, part)


def _add_heading_with_dates(
    document: Any,
    spec: dict[str, Any],
    part_name: str,
    left_text: str,
    dates: str,
    margins_mm: dict[str, Any],
    page_width_mm: float,
) -> Any:
    """左粗右灰的条目行：左侧机构+职位加粗，日期经右对齐 tab 停靠到版心右缘。"""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Mm, Pt

    part = spec["parts"][part_name]
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.alignment = _align("left")
    fmt.space_before = Pt(float(part.get("space_before_pt", 0)))
    fmt.space_after = Pt(float(part.get("space_after_pt", 2)))
    content_width = page_width_mm - float(margins_mm["left"]) - float(margins_mm["right"])
    fmt.tab_stops.add_tab_stop(Mm(content_width), WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run(left_text)
    _apply_run_style(run, spec, part)
    if dates:
        dates_part = dict(part)
        dates_part["bold"] = False
        dates_part["color"] = str(part.get("dates_color", "muted"))
        dates_part["size_pt"] = float(part.get("dates_size_pt", part.get("size_pt", 9)))
        run_dates = paragraph.add_run(f"\t{dates}")
        _apply_run_style(run_dates, spec, dates_part)
    return paragraph


def _set_hairline_below(paragraph: Any, spec: dict[str, Any], color_name: str = "hairline", width_pt: float = 0.5) -> None:
    """给段落加下边框发丝线（OXML）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    palette = spec["palette"]
    color = str(palette.get(color_name, color_name)).lstrip("#").upper()
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(int(width_pt * 8)))  # docx 单位为 1/8 pt
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_part_paragraph(
    document: Any,
    spec: dict[str, Any],
    part_name: str,
    text: str,
    *,
    style: str | None = None,
) -> Any:
    from docx.shared import Pt

    part = spec["parts"][part_name]
    paragraph = document.add_paragraph(style=style)
    fmt = paragraph.paragraph_format
    fmt.alignment = _align(str(part.get("align", "left")))
    fmt.space_before = Pt(float(part.get("space_before_pt", 0)))
    fmt.space_after = Pt(float(part.get("space_after_pt", 2)))
    if part.get("leading"):
        fmt.line_spacing = float(part["leading"])
    run = paragraph.add_run(text)
    _apply_run_style(run, spec, part)
    if part.get("rule") == "hairline-below":
        _set_hairline_below(paragraph, spec, str(part.get("rule_color", "hairline")), float(part.get("rule_width_pt", 0.5)))
        fmt.keep_with_next = True
    return paragraph


def render_docx(
    content: dict[str, Any],
    spec: dict[str, Any],
    path: Path,
    photo: Path | None = None,
) -> None:
    """按模板 spec 把底稿渲染为 DOCX（工作物，用户可直接编辑）。

    content 为现有 portable 底稿 schema（person/target/summary/skills/
    experience_sections/education）。photo 仅在 spec.photo != "none" 时生效。
    """
    from docx import Document
    from docx.shared import Mm, Pt

    document = Document()
    page = spec["page"]
    page_width_mm, page_height_mm = _page_dimensions_mm(page)
    margins = page["margins_mm"]
    section = document.sections[0]
    section.page_width = Mm(page_width_mm)
    section.page_height = Mm(page_height_mm)
    section.top_margin = Mm(float(margins["top"]))
    section.bottom_margin = Mm(float(margins["bottom"]))
    section.left_margin = Mm(float(margins["left"]))
    section.right_margin = Mm(float(margins["right"]))

    person = dict(content.get("person") or {})
    contact_sep = str(spec["parts"]["contact"].get("separator", "  |  "))
    contacts = [
        str(value)
        for value in (
            person.get("gender"),
            _format_age(person.get("age")),
            person.get("phone"),
            person.get("email"),
            person.get("city"),
            person.get("linkedin"),
            person.get("website"),
        )
        if value
    ]

    # —— 头部 ——
    if spec.get("photo") == "none" or photo is None:
        _add_part_paragraph(document, spec, "name", str(person.get("name") or ""))
        if contacts:
            _add_part_paragraph(document, spec, "contact", contact_sep.join(contacts))
    else:
        from docx.enum.table import WD_ALIGN_VERTICAL

        photo_cols = spec.get("photo_columns_mm") or [146, 36]
        photo_width = float(spec.get("photo_width_mm", 26))
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        left_cell, right_cell = table.cell(0, 0), table.cell(0, 1)
        left_cell.width = Mm(float(photo_cols[0]))
        right_cell.width = Mm(float(photo_cols[1]))
        name_part = spec["parts"]["name"]
        p_name = left_cell.paragraphs[0]
        p_name.paragraph_format.space_after = Pt(float(name_part.get("space_after_pt", 2)))
        run = p_name.add_run(str(person.get("name") or ""))
        _apply_run_style(run, spec, name_part)
        if contacts:
            contact_part = spec["parts"]["contact"]
            p_contact = left_cell.add_paragraph()
            p_contact.paragraph_format.space_after = Pt(float(contact_part.get("space_after_pt", 2)))
            run = p_contact.add_run(contact_sep.join(contacts))
            _apply_run_style(run, spec, contact_part)
        right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_photo = right_cell.paragraphs[0]
        p_photo.paragraph_format.space_after = Pt(0)
        p_photo.add_run().add_picture(str(photo), width=Mm(photo_width))

    header_rule = spec["parts"].get("header_rule")
    if header_rule:
        rule_paragraph = document.add_paragraph()
        rule_paragraph.paragraph_format.space_before = Pt(0)
        rule_paragraph.paragraph_format.space_after = Pt(float(header_rule.get("space_after_pt", 6)))
        _set_hairline_below(
            rule_paragraph,
            spec,
            color_name=str(header_rule.get("color", "hairline")),
            width_pt=float(header_rule.get("width_pt", 0.5)),
        )

    # —— 岗位目标 ——
    # match_score 属内部排序数据，不写入对外投递的简历
    target = dict(content.get("target") or {})
    target_line = format_resume_target(target)
    if target_line:
        _add_part_paragraph(document, spec, "target", target_line)

    # —— 摘要 ——
    summary = str(content.get("summary") or "").strip()
    summary_part = spec["parts"]["summary"]
    summary_layout = str(summary_part.get("layout", "paragraph"))
    if summary and summary_layout in {"numbered-section", "bullet-section"}:
        _add_part_paragraph(
            document, spec, "section_heading", str(summary_part.get("label") or "个人优势")
        )
        items = [piece.strip() for piece in re.split(r"[；;。]", summary) if piece.strip()]
        if summary_layout == "numbered-section":
            _add_numbered_items(document, spec, summary_part, items)
        else:
            _add_bulleted_items(document, spec, summary_part, items)
    elif summary:
        _add_part_paragraph(document, spec, "summary", summary)

    # —— 核心能力 ——
    skills = [str(item) for item in list(content.get("skills", []))]
    if skills and bool(spec["parts"].get("skills", {}).get("enabled", True)):
        _add_part_paragraph(document, spec, "skills_label", "核心能力")
        skills_sep = str(spec["parts"]["skills"].get("separator", " · "))
        _add_part_paragraph(document, spec, "skills", skills_sep.join(skills))

    # —— 经历模块 ——
    entry_part = spec["parts"]["entry_heading"]
    entry_sep = str(entry_part.get("separator", " · "))
    bullet_part = spec["parts"]["bullet"]
    numbered = str(bullet_part.get("list", "bullet")) == "number"
    for section_payload in list(content.get("experience_sections", [])):
        section_payload = dict(section_payload)
        _add_part_paragraph(
            document, spec, "section_heading", resume_section_title(section_payload)
        )
        for entry_payload in list(section_payload.get("entries", [])):
            entry_payload = dict(entry_payload)
            if str(entry_part.get("dates_align", "inline")) == "right":
                left_text = "  ".join(
                    value
                    for value in (
                        str(entry_payload.get("organization") or ""),
                        str(entry_payload.get("role") or ""),
                    )
                    if value
                )
                if left_text:
                    _add_heading_with_dates(
                        document, spec, "entry_heading",
                        left_text, str(entry_payload.get("dates") or ""), margins,
                        page_width_mm,
                    )
            else:
                heading = entry_sep.join(
                    value
                    for value in (
                        str(entry_payload.get("organization") or ""),
                        str(entry_payload.get("role") or ""),
                        str(entry_payload.get("dates") or ""),
                    )
                    if value
                )
                if heading:
                    _add_part_paragraph(document, spec, "entry_heading", heading)
            grouped: dict[str, list[str]] = {}
            section_title = str(section_payload.get("title") or "")
            is_work = "工作" in section_title or "实习" in section_title
            default_label = "工作内容" if is_work else "项目内容"
            for bullet_payload in list(entry_payload.get("bullets", [])):
                bullet = dict(bullet_payload)
                text = str(bullet.get("text") or "").strip()
                if not text:
                    continue
                label = str(bullet.get("label") or default_label).strip()
                if label not in {"工作内容", "工作业绩", "项目内容", "项目业绩"}:
                    label = default_label
                grouped.setdefault(label, []).append(text)
            label_order = (
                ("工作内容", "工作业绩")
                if is_work
                else ("项目内容", "项目业绩")
            )
            for label in label_order:
                bullets = grouped.get(label, [])
                if not bullets:
                    continue
                if "entry_label" in spec["parts"]:
                    _add_part_paragraph(document, spec, "entry_label", f"{label}：")
                if numbered:
                    _add_numbered_items(document, spec, bullet_part, bullets)
                else:
                    _add_bulleted_items(document, spec, bullet_part, bullets)

    # —— 教育经历 ——
    education = list(content.get("education", []))
    if education:
        _add_part_paragraph(document, spec, "section_heading", "教育经历")
        edu_part = spec["parts"]["education_item"]
        edu_sep = str(edu_part.get("separator", " · "))
        for item in education:
            item = dict(item)
            if str(edu_part.get("dates_align", "inline")) == "right":
                left_text = "  ".join(
                    value
                    for value in (
                        str(item.get("institution") or ""),
                        str(item.get("major") or ""),
                        str(item.get("degree") or ""),
                    )
                    if value
                )
                start, end = str(item.get("start") or ""), str(item.get("end") or "")
                dates = f"{start} - {end}" if start and end else (end or start)
                if left_text:
                    _add_heading_with_dates(
                        document, spec, "education_item", left_text, dates, margins,
                        page_width_mm,
                    )
                coursework = [str(course) for course in list(item.get("coursework") or []) if course]
                if coursework:
                    add_items = _add_numbered_items if numbered else _add_bulleted_items
                    add_items(document, spec, bullet_part, [f"主修课程：{'、'.join(coursework)}"])
            else:
                label = edu_sep.join(
                    value
                    for value in (
                        str(item.get("institution") or ""),
                        str(item.get("major") or ""),
                        str(item.get("degree") or ""),
                        str(item.get("end") or ""),
                    )
                    if value
                )
                if label:
                    _add_part_paragraph(document, spec, "education_item", label)

    evaluation = str(content.get("self_evaluation") or "").strip()
    if evaluation:
        _add_part_paragraph(document, spec, "section_heading", "自我评价")
        _add_part_paragraph(document, spec, "summary", evaluation)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
